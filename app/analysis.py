"""Resume parsing and scoring. Deterministic: the same resume always gets the same result."""
import csv
import io
import json
import re
from functools import lru_cache
from pathlib import Path

import numpy as np

from app.embed import embed

DATA = Path(__file__).resolve().parent.parent / "data"


# ---------- parsing ----------

def parse_resume(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    elif ext == "docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts += [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
        text = "\n".join(parts)
    else:
        raise ValueError("Upload a PDF or DOCX file.")
    if len(text.split()) < 50:
        raise ValueError(
            "Couldn't read enough text from this file. If it's a scanned PDF, "
            "export it as a text PDF from Word or Google Docs."
        )
    return text


# ---------- skills ----------

@lru_cache
def _skill_patterns():
    """data/skills.csv: canonical name, then aliases. A leading '=' (or an all-caps
    acronym) means case-sensitive, so 'Excel' the tool doesn't match 'excel at'."""
    out = []
    with open(DATA / "skills.csv") as f:
        for row in csv.reader(f):
            terms = [t.strip() for t in row if t.strip()]
            if not terms:
                continue
            pats = []
            for t in terms:
                sensitive = t.startswith("=") or (t.isupper() and len(t) <= 5)
                t = t.lstrip("=")
                flags = 0 if sensitive else re.I
                pats.append(re.compile(rf"(?<![\w+#.-]){re.escape(t)}(?![\w+#&-])", flags))
            out.append((terms[0].lstrip("="), pats))
    return out


URL = re.compile(r"\S+\.(?:com|in|io|org|dev|me|net)\S*", re.I)


def extract_skills(text: str) -> list[str]:
    text = URL.sub(" ", text)  # github.com/x is a link, not a GitHub skill
    return [name for name, pats in _skill_patterns() if any(p.search(text) for p in pats)]


def skill_evidence(text: str, skills: list[str]) -> dict:
    """Resume lines that back up each skill, preferring prose bullets over comma lists."""
    pats = dict(_skill_patterns())
    lines = [ln.strip(" \t•●▪◦*-") for ln in text.splitlines()]
    lines = [ln for ln in lines if len(ln) > 30]
    out = {}
    for s in skills:
        hits = [ln for ln in lines if any(p.search(ln) for p in pats[s])]
        hits.sort(key=lambda ln: ln.count(","))
        if hits:
            out[s] = hits[:2]
    return out


# ---------- profile / structure ----------

EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE = re.compile(r"\+?\d[\d\s()-]{8,16}\d")
LINK = re.compile(r"(?:https?://)?(?:www\.)?(?:linkedin\.com/in|github\.com)/[\w-]+/?", re.I)
SECTIONS = {
    "Education": r"education|academics?",
    "Experience": r"experience|internships?|employment|work history",
    "Projects": r"projects?",
    "Skills": r"skills|technologies|tech stack",
}


def profile(text: str) -> dict:
    phones = [p for p in PHONE.findall(text) if 10 <= len(re.sub(r"\D", "", p)) <= 13]
    name = next(
        (ln.strip() for ln in text.splitlines()[:5]
         if re.fullmatch(r"[A-Za-z][A-Za-z .'-]{2,40}", ln.strip()) and 2 <= len(ln.split()) <= 4),
        "",
    )
    email = EMAIL.search(text)
    return {
        "name": name.title() if name.isupper() else name,
        "email": email.group(0) if email else "",
        "phone": phones[0].strip() if phones else "",
        "links": list(dict.fromkeys(m.group(0) for m in LINK.finditer(text))),
    }


def _quantified_lines(text: str) -> int:
    # ponytail: "has a number that isn't a year" heuristic; misses spelled-out numbers.
    n = 0
    for ln in URL.sub(" ", text).splitlines():
        if len(ln) < 30 or re.search(r"cgpa|gpa|percentage|\bclass\b|@|\+91", ln, re.I):
            continue
        if re.search(r"(?<![\w/.-])(?!(?:19|20)\d\d(?!\d))\d+(?:\.\d+)?\s*(?:%|x\b|\+|k\b)?", ln):
            n += 1
    return n


def ats_report(text: str, skills: list[str], target: list[str], target_label: str) -> dict:
    checks = []

    def add(label, points, max_points, tip):
        points = round(points)
        checks.append({"label": label, "points": points, "max": max_points,
                       "tip": tip if points < max_points else ""})

    heads = [ln.strip().lower() for ln in text.splitlines()
             if 0 < len(ln.strip()) <= 40 and len(ln.split()) <= 5]
    for name, pat in SECTIONS.items():
        found = any(re.search(rf"\b(?:{pat})\b", h) for h in heads)
        add(f"{name} section", 6 if found else 0, 6,
            f"Add a clear '{name}' heading so ATS parsers can find this section.")

    p = profile(text)
    add("Email", 4 if p["email"] else 0, 4, "Put your email address in the header.")
    add("Phone", 4 if p["phone"] else 0, 4, "Put a phone number in the header.")
    add("LinkedIn / GitHub", 4 if p["links"] else 0, 4, "Link your LinkedIn and GitHub profiles in the header.")

    q = _quantified_lines(text)
    add("Measurable impact", min(q, 6) / 6 * 20, 20,
        f"{q} line(s) include a number. Aim for 6 or more bullets with metrics (accuracy, users, time saved).")

    words = len(text.split())
    length_pts = 10 if 350 <= words <= 900 else 6 if 250 <= words <= 1200 else 2
    add("Length", length_pts, 10,
        f"{words} words. One page of 350 to 900 words reads best." if words < 350 or words > 900 else "")

    add("Skills listed", min(len(skills), 12) / 12 * 10, 10,
        "List your tools and technologies explicitly; ATS filters match exact names.")

    have = set(skills)
    missing = [s for s in target if s not in have]
    cov = (len(target) - len(missing)) / len(target) if target else 1
    add(f"Keywords ({target_label})", cov * 24, 24,
        "Missing: " + ", ".join(missing[:6]) + ". Add the ones you genuinely have." if missing else "")

    return {"score": sum(c["points"] for c in checks), "checks": checks}


# ---------- semantic matching ----------

def _chunks(text: str, words: int = 60) -> list[str]:
    """MiniLM truncates long inputs, so embed the resume in ~60-word pieces instead of all at once."""
    lines = [ln.strip() for ln in text.splitlines() if len(ln.strip()) > 20]
    chunks, cur = [], []
    for ln in lines:
        cur.append(ln)
        if sum(len(x.split()) for x in cur) >= words:
            chunks.append(" ".join(cur))
            cur = []
    if cur:
        chunks.append(" ".join(cur))
    return chunks or [text[:1500]]


def text_vector(text: str, extra: list[str] = ()) -> np.ndarray:
    v = embed(_chunks(text) + [", ".join(extra)] if extra else _chunks(text)).mean(0)
    return v / np.linalg.norm(v)


def scale(sim: float, lo: float = 0.2, hi: float = 0.75) -> float:
    """Map raw MiniLM cosine onto 0..1. Calibrated on real resume/JD pairs:
    an unrelated JD scores ~0.18, a strong match ~0.68."""
    return float(min(max((sim - lo) / (hi - lo), 0.0), 1.0))


@lru_cache
def _roles():
    roles = json.loads((DATA / "roles.json").read_text())
    vecs = embed([f"{r}. Skills: {', '.join(s)}" for r, s in roles.items()])
    return roles, list(roles), vecs


def predict_roles(resume_vec: np.ndarray, skills: list[str], top: int = 5) -> list[dict]:
    roles, names, vecs = _roles()
    sims = vecs @ resume_vec
    have = set(skills)
    out = []
    for name, sim in zip(names, sims):
        req = roles[name]
        cov = sum(s in have for s in req) / len(req)
        out.append({
            "role": name,
            "fit": round(100 * (0.5 * cov + 0.5 * scale(sim))),
            "have": [s for s in req if s in have],
            "missing": [s for s in req if s not in have],
        })
    return sorted(out, key=lambda r: -r["fit"])[:top]


def analyze(text: str, job_description: str = "") -> dict:
    skills = extract_skills(text)
    rv = text_vector(text, skills)
    roles = predict_roles(rv, skills)

    match = None
    jd = job_description.strip()
    if jd:
        job_skills = extract_skills(jd)
        sem = scale(float(rv @ text_vector(jd, job_skills)))
        matched = [s for s in job_skills if s in skills]
        cov = len(matched) / len(job_skills) if job_skills else None
        score = 0.6 * cov + 0.4 * sem if cov is not None else sem
        match = {
            "score": round(100 * score),
            "semantic": round(100 * sem),
            "coverage": None if cov is None else round(100 * cov),
            "matched": matched,
            "missing": [s for s in job_skills if s not in skills],
        }
        target, label = job_skills, "job description"
    else:
        target, label = roles[0]["have"] + roles[0]["missing"], roles[0]["role"]

    return {
        "profile": profile(text),
        "skills": skills,
        "roles": roles,
        "match": match,
        "ats": ats_report(text, skills, target, label),
        "evidence": skill_evidence(text, skills),
    }
